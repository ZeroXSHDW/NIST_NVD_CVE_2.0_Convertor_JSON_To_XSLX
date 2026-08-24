"""Regression tests for the fail-closed workbook integrity boundary."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from json_to_xlsx import build_workbook
from validate_xlsx import main, validate_workbook


def _write_feed(data_dir: Path, vulnerabilities: list[dict]) -> Path:
    data_dir.mkdir()
    feed = data_dir / "nvdcve-2.0-2024.json"
    feed.write_text(json.dumps({"vulnerabilities": vulnerabilities}), encoding="utf-8")
    return feed


def _cve(cve_id: str) -> dict:
    return {
        "cve": {
            "id": cve_id,
            "descriptions": [{"lang": "en", "value": f"Description for {cve_id}"}],
            "metrics": {},
        }
    }


def test_validator_accepts_complete_workbook_and_malformed_record(tmp_path):
    data_dir = tmp_path / "feeds"
    _write_feed(data_dir, [{}, {"cve": None}, _cve("CVE-2024-0001")])
    output = tmp_path / "compiled.xlsx"

    build_workbook(data_dir=data_dir, output_file=output)

    assert validate_workbook(data_dir=data_dir, xlsx_path=output) == []
    assert main(data_dir=data_dir, xlsx_path=output) == 0


def test_validator_rejects_truncated_annual_sheet_and_returns_nonzero(tmp_path):
    data_dir = tmp_path / "feeds"
    _write_feed(data_dir, [_cve("CVE-2024-0001"), _cve("CVE-2024-0002")])
    output = tmp_path / "compiled.xlsx"
    build_workbook(data_dir=data_dir, output_file=output)

    workbook = load_workbook(output)
    try:
        workbook["2024"].delete_rows(2)
        workbook.save(output)
    finally:
        workbook.close()

    mismatches = validate_workbook(data_dir=data_dir, xlsx_path=output)
    assert any(item[:2] == ("2024", "row_count") for item in mismatches)
    assert main(data_dir=data_dir, xlsx_path=output) == 1


def test_validator_rejects_truncated_index(tmp_path):
    data_dir = tmp_path / "feeds"
    _write_feed(data_dir, [_cve("CVE-2024-0001"), _cve("CVE-2024-0002")])
    output = tmp_path / "compiled.xlsx"
    build_workbook(data_dir=data_dir, output_file=output)

    workbook = load_workbook(output)
    try:
        workbook["INDEX"].delete_rows(2)
        workbook.save(output)
    finally:
        workbook.close()

    mismatches = validate_workbook(data_dir=data_dir, xlsx_path=output)
    assert any(item[:2] == ("INDEX", "row_count") for item in mismatches)
    assert main(data_dir=data_dir, xlsx_path=output) == 1
